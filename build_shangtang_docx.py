import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("shangtang-api-test-report.md")
OUTPUT = Path("shangtang-api-test-report.docx")
PAGE_WIDTH = Inches(7.1)
FONT = "Microsoft YaHei"
MONO = "Consolas"


def set_run_font(run, name=FONT, size=None, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    props.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))")


def add_inline(paragraph, text, size=None, default_bold=False):
    position = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, bold=default_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO, size or 9, bold=default_bold, color="17365D")
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EEF3F8")
            run._element.get_or_add_rPr().append(shading)
        else:
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, bold=default_bold)


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=90, bottom=75, end=90):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def remove_table_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    props.append(borders)


def prevent_row_split(row):
    props = row._tr.get_or_add_trPr()
    if props.find(qn("w:cantSplit")) is None:
        props.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def set_fixed_layout(table):
    props = table._tbl.tblPr
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")


def table_widths(rows, total_width):
    count = len(rows[0])
    weights = []
    for col in range(count):
        values = [len(re.sub(r"[`*]", "", row[col])) for row in rows if col < len(row)]
        weights.append(max(8, min(max(values, default=8), 52)))
    total = sum(weights)
    return [int(total_width * weight / total) for weight in weights]


def make_table(document, rows):
    widths = table_widths(rows, PAGE_WIDTH)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_fixed_layout(table)
    set_table_borders(table)
    for row_index, row_data in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_header(row)
        for col_index, value in enumerate(row_data):
            table.cell(row_index, col_index).width = widths[col_index]
            cell_item = table.cell(row_index, col_index)
            cell_item.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell_item)
            set_cell_shading(cell_item, "1F4E79" if row_index == 0 else ("F2F6FA" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell_item.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) < 22 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value.strip(), size=8.2, default_bold=row_index == 0)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    return table


def add_nonbreaking_table(document, rows, continuation=False):
    document.add_page_break()
    if continuation:
        label = document.add_paragraph()
        label.paragraph_format.keep_with_next = True
        label.paragraph_format.space_before = Pt(4)
        label.paragraph_format.space_after = Pt(4)
        run = label.add_run("表格续")
        set_run_font(run, size=9, bold=True, color="666666")
    make_table(document, rows)


def split_table(rows):
    header, body = rows[0], rows[1:]
    longest = max((len(cell) for row in body for cell in row), default=0)
    columns = len(header)
    if longest > 260:
        limit = 2
    elif columns >= 6:
        limit = 6
    elif columns >= 5:
        limit = 7
    else:
        limit = 9
    return [[header] + body[index:index + limit] for index in range(0, len(body), limit)] or [[header]]


def clean_heading(text):
    if text.startswith("SenseAudio"):
        return "SenseAudio 文本 图像 视频 API 详细测试报告"
    text = re.sub(r"[：:，,。.!！？?（）()\[\]【】]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def add_heading(document, text, level, first_heading):
    style = "Title" if level == 1 and first_heading else f"Heading {min(level, 3)}"
    paragraph = document.add_paragraph(style=style)
    if level == 1 and not first_heading:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level > 1 else 0)
    paragraph.paragraph_format.space_after = Pt(7 if level <= 2 else 5)
    add_inline(paragraph, clean_heading(text), default_bold=True)
    return paragraph


def add_code_block(document, lines):
    for index, line in enumerate(lines or [""]):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.08)
        paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(4 if index == len(lines) - 1 else 0)
        paragraph.paragraph_format.line_spacing = 1.0
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F3F5F7")
        paragraph._p.get_or_add_pPr().append(shading)
        run = paragraph.add_run(line)
        set_run_font(run, MONO, size=8.3, color="202A35")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("SenseAudio API 测试报告  第 ")
    set_run_font(run, size=8, color="777777")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8, color="777777")


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.28)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    add_page_number(section.footer.paragraphs[0])
    document.core_properties.title = "SenseAudio 文本 图像 视频 API 详细测试报告"
    document.core_properties.subject = "文本 图片 视频模型接口与错误码测试"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def convert(source, output):
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    first_heading = True
    intro_added = False
    index = 0

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            add_heading(document, heading.group(2), level, first_heading)
            if first_heading:
                first_heading = False
                intro = document.add_paragraph()
                intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
                intro.paragraph_format.space_after = Pt(14)
                add_inline(intro, "远程调用实测结果 请求与响应契约 状态码 错误码和页面验收标准", size=10.5)
                intro_added = True
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            add_code_block(document, block)
            continue
        if line.lstrip().startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index + 1]):
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in table_lines]
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            for chunk_index, chunk in enumerate(split_table(rows)):
                add_nonbreaking_table(document, chunk, continuation=chunk_index > 0)
            continue
        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            next_line = lines[index]
            if not next_line.strip() or next_line.startswith("#") or next_line.startswith("```") or next_line.lstrip().startswith("|") or re.match(r"^\s*(\d+\.|[-*])\s+", next_line):
                break
            paragraph_lines.append(next_line.strip())
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else OUTPUT
    convert(source, output)
    print(output.resolve())
